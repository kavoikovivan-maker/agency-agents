"""Upload validation and safe text extraction for attached files."""
from __future__ import annotations

from pathlib import Path

from fastapi import HTTPException, UploadFile

from .config import settings

ALLOWED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".xlsx", ".xls", ".docx"}


def validate_upload(filename: str, size: int) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext or 'unknown'}")
    max_bytes = settings.max_upload_mb * 1024 * 1024
    if size > max_bytes:
        raise HTTPException(status_code=400, detail=f"File exceeds {settings.max_upload_mb}MB limit")


def extract_text(path: Path) -> tuple[str, str]:
    """Best-effort text extraction. Returns (text, error). Never raises."""
    ext = path.suffix.lower()
    try:
        if ext in {".txt", ".md", ".csv", ".json"}:
            return path.read_text(encoding="utf-8", errors="ignore"), ""
        if ext == ".pdf":
            return _extract_pdf(path), ""
        if ext == ".docx":
            return _extract_docx(path), ""
        if ext in {".xlsx", ".xls"}:
            return _extract_xlsx(path), ""
        return "", f"No extractor for {ext}"
    except Exception as exc:  # noqa: BLE001 - extraction must never crash the upload flow
        return "", f"Extraction failed: {exc}"


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


async def save_upload(upload: UploadFile, dest: Path) -> int:
    size = 0
    with dest.open("wb") as f:
        while chunk := await upload.read(1024 * 1024):
            size += len(chunk)
            if size > settings.max_upload_mb * 1024 * 1024:
                dest.unlink(missing_ok=True)
                raise HTTPException(status_code=400, detail=f"File exceeds {settings.max_upload_mb}MB limit")
            f.write(chunk)
    return size
